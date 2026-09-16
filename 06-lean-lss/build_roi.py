"""
產生 roi-validation.xlsx —— 透明的 ROI 效益驗證模型。

原則（見 docs/superpowers/specs 的第 3 節）：
  綠 = 實測值（親自量測，不可改）
  黃 = 假設值（可自由替換；換成 ASML 真實數字後，公式立刻重算）
  藍 = 計算值（Excel 公式，不寫死）

所有金額用真實 Excel 公式表達，方便面試時當場改假設、看 ROI 變化。

執行：py -3 build_roi.py
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = Path(__file__).resolve().parent / "roi-validation.xlsx"

# ---- 樣式 ----
GREEN = PatternFill("solid", fgColor="C6EFCE")   # 實測
YELLOW = PatternFill("solid", fgColor="FFEB9C")  # 假設
BLUE = PatternFill("solid", fgColor="DDEBF7")    # 計算
HEAD = PatternFill("solid", fgColor="1F4E78")    # 區塊標題
SUBHEAD = PatternFill("solid", fgColor="8EA9DB")

WHITE_BOLD = Font(color="FFFFFF", bold=True, size=12)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NT = '#,##0'
PCT1 = '0.0%'


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "ROI"
    ws.sheet_view.showGridLines = False

    widths = {"A": 30, "B": 16, "C": 12, "D": 40}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w

    def row(r, a, b=None, c=None, d=None, fill=None, num=None, bold=False):
        ws.cell(r, 1, a)
        if bold:
            ws.cell(r, 1).font = BOLD
        if b is not None:
            cell = ws.cell(r, 2, b)
            if fill:
                cell.fill = fill
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="right")
            if num:
                cell.number_format = num
        if c is not None:
            ws.cell(r, 3, c).alignment = Alignment(horizontal="center")
        if d is not None:
            ws.cell(r, 4, d).font = Font(size=9, color="595959")

    def block(r, title):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        cell = ws.cell(r, 1, title)
        cell.fill = HEAD
        cell.font = WHITE_BOLD

    # ===== 標題與圖例 =====
    ws.merge_cells("A1:D1")
    ws["A1"] = "ROI 效益驗證 — 預測維護資料流程自動化"
    ws["A1"].font = Font(bold=True, size=15)
    ws.merge_cells("A2:D2")
    ws["A2"] = ("誠信原則：綠=實測值(不可改) · 黃=假設值(可替換) · 藍=Excel公式(自動重算)。"
                "把黃色換成 ASML 真實數字，下方 ROI 立刻重算。")
    ws["A2"].font = Font(size=10, italic=True, color="595959")

    ws["A3"] = "實測"; ws["A3"].fill = GREEN; ws["A3"].border = BORDER
    ws["B3"] = "假設"; ws["B3"].fill = YELLOW; ws["B3"].border = BORDER
    ws["C3"] = "計算"; ws["C3"].fill = BLUE; ws["C3"].border = BORDER

    # ===== 區塊 A：流程自動化效益（第二層，主角）=====
    block(5, "A. 流程自動化效益（第二層 · 消除人工資料處理浪費）")
    ws.cell(6, 1, "項目").font = BOLD
    ws.cell(6, 2, "值").font = BOLD
    ws.cell(6, 3, "類型").font = BOLD
    ws.cell(6, 4, "來源 / 說明").font = BOLD

    row(7,  "單次手動耗時（秒）", 900, "實測🟢", "碼錶實測：手動清理彙整一次 = 15 分鐘", GREEN)
    row(8,  "單次自動耗時（秒）", 1.6, "實測🟢", "VBA 巨集 Timer 回報", GREEN, num="0.0")
    row(9,  "單次省時（秒）", "=B7-B8", "計算", "手動 − 自動", BLUE, num="0.0")
    row(10, "每天執行次數", 2, "假設🟡", "每日跑幾次報表（可改）", YELLOW)
    row(11, "使用人數", 1, "假設🟡", "幾位工程師做同樣作業（可改）", YELLOW)
    row(12, "年工作天", 240, "假設🟡", "一年工作天數（可改）", YELLOW)
    row(13, "年省工時（小時）", "=B9*B10*B11*B12/3600", "計算", "省時×次數×人數×天數÷3600", BLUE, num="#,##0.0")
    row(14, "人力成本（NT$/小時）", 500, "假設🟡", "工程師時薪（可改）", YELLOW)
    row(15, "年省人力成本（NT$）", "=B13*B14", "計算", "年省工時 × 時薪", BLUE, num=NT)
    row(16, "開發投入工時（小時）", 40, "假設🟡", "一次性全導入工時：需求釐清、巨集開發(AI輔助約8hr)、測試、Power Automate串接、看板、文件與推廣訓練", YELLOW)
    row(17, "導入成本（NT$）", "=B16*B14", "計算", "開發工時 × 時薪", BLUE, num=NT)
    row(18, "回收期（工作天）", "=B17/(B15/B12)", "計算", "導入成本 ÷ 每日省下的成本", BLUE, num="0.0")
    row(19, "首年淨效益（NT$）", "=B15-B17", "計算", "年省 − 導入成本", BLUE, num=NT)
    row(20, "首年 ROI", "=(B15-B17)/B17", "計算", "(年省 − 導入) ÷ 導入", BLUE, num="0%")
    ws.cell(20, 1).font = BOLD
    ws.cell(20, 2).font = Font(bold=True, color="C00000")

    # ===== 區塊 B：品質效益（不會報錯的錯誤）=====
    block(23, "B. 品質效益（消除靜默錯誤 · Defects）")
    row(24, "狀態大小寫需正規化（筆）", 2001, "實測🟢", "巨集統計：Warning/warning/WARNING 混用", GREEN, num="#,##0")
    row(25, "機台代號帶空白比例", 0.15, "實測🟢", "資料特性；手動漏 TRIM 會低估警告數", GREEN, num=PCT1)
    ws.merge_cells("A26:D26")
    ws["A26"] = ("手動樞紐若漏 TRIM 或未統一大小寫，會把同一機台/同一 KPI 拆成多列多欄，"
                 "報表看似完成卻默默少算——且 Excel 不報錯。巨集規則固定，缺陷率 = 0。")
    ws["A26"].font = Font(size=9, color="595959")
    ws["A26"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[26].height = 42

    # ===== 區塊 C：AI 模型效益（第一層 · 技術亮點）=====
    block(29, "C. AI 模型效益（第一層 · 設備健康判讀）")
    row(30, "模型 Accuracy", 0.85, "實測🟢", "VAE + 最近 support set，100 台最後一個 cycle", GREEN, num=PCT1)
    row(31, "Warning 的 Recall", 0.939, "實測🟢", "33 台快壞的，抓到 31 台", GREEN, num=PCT1)
    row(32, "漏判台數", 2, "實測🟢", "該預警卻沒抓到（代價最高的錯誤）", GREEN, num="0")
    row(33, "維護成本降低", 0.84, "假設🟡", "情境假設：公開資料集無真實成本，非實測", YELLOW, num=PCT1)
    row(34, "停機時間降低", 0.36, "假設🟡", "情境假設：同上", YELLOW, num=PCT1)

    # ===== 底部：兩層效益分列聲明 =====
    ws.merge_cells("A36:D36")
    ws["A36"] = ("兩層效益分列：A/B 為本案主角（流程自動化，效益可實測）；C 為技術亮點"
                 "（設備健康，公開資料集下的情境假設）。未混算、未灌水。")
    ws["A36"].font = Font(size=10, italic=True, bold=True, color="1F4E78")
    ws["A36"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[36].height = 30

    wb.save(OUT)
    print(f"已產生：{OUT}")


if __name__ == "__main__":
    main()
